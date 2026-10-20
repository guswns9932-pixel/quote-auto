import hashlib
import json
import logging
import os
import sys
import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "no_com"))

from openpyxl import Workbook, load_workbook
from core import QuoteState, SheetName
from document_work import publish_document
from excel_io import generate_quote, generate_quote_multi
from package_local import package_local


class LocalWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.local = self.root / "local"
        # Override actual storage roots, never touch developer/user documents.
        for module in ("document_work", "app_paths"):
            mock = patch(module + ".local_data_dir", return_value=self.local)
            mock.start()
            self.addCleanup(mock.stop)
        self.cloud = self.root / "cloud"
        self.cloud.mkdir()
        self.template = self.cloud / "통합 양식.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        for name in SheetName.REQUIRED:
            wb.create_sheet(name)
        wb.save(self.template)
        wb.close()
        self.original = self.template.read_bytes()
        self.state = QuoteState(template_path=str(self.template),
                                output_root=str(self.cloud / "결과"))

    def test_three_quote_types_publish_valid_workbooks_and_keep_source(self):
        for kind, sheet in (("국내", SheetName.SPEC), ("중국", SheetName.QUOTE_CN),
                            ("미국", SheetName.QUOTE_US)):
            with self.subTest(kind=kind):
                path = Path(generate_quote(self.state, kind, [], {"D": "PR1", "H": 1}))
                self.assertTrue(path.is_relative_to(self.cloud / "결과"))
                wb = load_workbook(path)
                self.assertEqual(wb[sheet].sheet_state, "visible")
                if kind == "국내":
                    self.assertIn(self.state.investor_name, wb[sheet]["B4"].value)
                wb.close()
        self.assertEqual(self.template.read_bytes(), self.original)
        self.assertEqual(self.state.template_path, str(self.template))
        self.assertEqual(list((self.local / "work").iterdir()), [])

    def test_batch_keeps_documents_distinct(self):
        rows = [{"D": "PR1", "E": "1"}, {"D": "PR1", "E": "2"}]
        result = generate_quote_multi(self.state, [], rows)
        self.assertEqual(len(result), 2)
        self.assertTrue(all(path and not error for _, path, error in result), result)
        self.assertNotEqual(result[0][1], result[1][1])

    def test_request_snapshot_preserves_selected_row_and_original(self):
        request = self.cloud / "의뢰.xlsx"
        wb = Workbook()
        wb.active.title = "requests"
        wb.active["D3"] = "PR-selected"
        wb.active["H3"] = 4
        wb.save(request); wb.close()
        original = request.read_bytes()
        self.state.request_path = str(request)
        self.state.request_sheet_name = "requests"
        path = generate_quote(self.state, "국내", [], {"row_idx": 3, "H": 4})
        wb = load_workbook(path)
        self.assertEqual(wb[SheetName.REQ_COPY]["D2"].value, "PR-selected")
        self.assertEqual(wb[SheetName.REQ_COPY]["H2"].value, 4)
        wb.close()
        self.assertEqual(request.read_bytes(), original)
        self.assertEqual(self.state.request_path, str(request))

    def test_batch_upload_errors_are_reported_with_recovery_paths(self):
        blocked = self.cloud / "blocked"
        blocked.write_text("keep")
        self.state.output_root = str(blocked)
        result = generate_quote_multi(self.state, [], [{"E": "1"}, {"E": "2"}])
        self.assertEqual(len(result), 2)
        self.assertTrue(all(not p and "완성된 로컬 파일" in err for _, p, err in result))
        self.assertEqual(len(list((self.local / "work").glob("*/output/**/*.xlsx"))), 2)

    def test_missing_input_cleans_staging_without_creating_results(self):
        self.state.template_path = str(self.cloud / "missing.xlsx")
        with self.assertRaises(FileNotFoundError):
            generate_quote(self.state, "국내", [], {})
        self.assertEqual(list((self.local / "work").iterdir()), [])
        self.assertFalse((self.cloud / "결과").exists())

    def test_logs_are_local_and_rotated(self):
        import app_boot
        root = logging.getLogger("QuoteApp")
        existing, level = root.handlers[:], root.level
        root.handlers = []
        try:
            with patch.object(app_boot.sys, "stdout", None), patch.object(app_boot, "_LOG_MAX_BYTES", 128):
                app_boot._setup_logging()
                for _ in range(10):
                    root.info("test local log rotation " * 3)
            logs = list((self.local / "logs").glob("*.log*"))
            self.assertEqual(len(logs), 2)
            self.assertTrue(all(p.stat().st_size < 512 for p in logs))
        finally:
            for handler in root.handlers:
                handler.close()
            root.handlers = existing
            root.setLevel(level)

    def test_upload_failure_retains_completed_local_document(self):
        # A file in place of the destination directory simulates inaccessible storage.
        blocked = self.cloud / "blocked"
        blocked.write_text("existing data")
        self.state.output_root = str(blocked)
        with self.assertRaisesRegex(OSError, "완성된 로컬 파일"):
            generate_quote(self.state, "국내", [], {})
        completed = list((self.local / "work").glob("*/output/**/*.xlsx"))
        self.assertEqual(len(completed), 1)
        wb = load_workbook(completed[0]); wb.close()
        self.assertEqual(blocked.read_text(), "existing data")

    def test_parallel_publish_never_overwrites(self):
        sources = []
        for i in range(12):
            source = self.root / f"source-{i}"
            source.write_text(str(i))
            sources.append(source)
        dest = self.cloud / "same.xlsx"
        dest.write_text("original")
        with ThreadPoolExecutor(max_workers=6) as pool:
            outputs = list(pool.map(lambda src: publish_document(src, dest), sources))
        self.assertEqual(len(set(outputs)), 12)
        self.assertEqual(dest.read_text(), "original")
        self.assertEqual({Path(p).read_text() for p in outputs}, {str(i) for i in range(12)})
        self.assertFalse(list(self.cloud.glob(".quote-upload-*")))

    def test_interrupted_copy_does_not_publish_partial_file(self):
        source = self.root / "source"
        source.write_text("completed document")
        dest = self.cloud / "output.xlsx"
        def fail_copy(src, dst):
            Path(dst).write_text("partial")
            raise OSError("connection lost")
        with patch("document_work.shutil.copyfile", side_effect=fail_copy):
            with self.assertRaises(OSError):
                publish_document(source, dest)
        self.assertFalse(dest.exists())
        self.assertFalse(list(self.cloud.glob(".quote-upload-*")))
        self.assertEqual(source.read_text(), "completed document")

    def test_package_contains_full_onedir_and_valid_checksum(self):
        build = self.root / "build"
        (build / "_internal").mkdir(parents=True)
        (build / "견적 자동화.exe").write_bytes(b"launcher fixture")
        (build / "_internal" / "runtime.dll").write_bytes(b"dependency fixture")
        dist = package_local(build, "견적 자동화")
        manifest = json.loads((dist / "release.json").read_text(encoding="utf-8"))
        archive = dist / manifest["package"]
        self.assertEqual(hashlib.sha256(archive.read_bytes()).hexdigest(), manifest["sha256"])
        with zipfile.ZipFile(archive) as z:
            self.assertEqual(z.read("_internal/runtime.dll"), b"dependency fixture")
            self.assertEqual(z.read(manifest["launcher"]), b"launcher fixture")
        self.assertTrue((dist / "Install.cmd").exists())
        # Repackaging doesn't mutate the build or remove previous release payloads.
        (build / "견적 자동화.exe").write_bytes(b"new launcher fixture")
        package_local(build, "견적 자동화")
        self.assertTrue(archive.exists())
        new = json.loads((dist / "release.json").read_text(encoding="utf-8"))
        self.assertNotEqual(new["sha256"], manifest["sha256"])

    def test_packaging_rejects_missing_launcher(self):
        with self.assertRaises(FileNotFoundError):
            package_local(self.root, "missing")


if __name__ == "__main__":
    unittest.main()
